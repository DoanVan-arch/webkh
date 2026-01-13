from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file
from flask_login import login_required, current_user
from models import (db, DeTai, TaiLieu, TiemLuc, SangKien, User, DON_VI_LIST, LOAI_TAI_LIEU, 
                    CHUC_DANH_LIST, NHA_GIAO_LIST, VAI_TRO_BAN, VAI_TRO_HOI_DONG,
                    BanDeTai, HoiDongDeCuongDeTai, HoiDongNghiemThuDeTai,
                    BanSangKien, HoiDongDeCuongSangKien, HoiDongNghiemThuSangKien,
                    BanTaiLieu, HoiDongDeCuongTaiLieu, HoiDongNghiemThuTaiLieu, CAP_DO_LIST)
from datetime import datetime
from sqlalchemy import func
from utils.file_handler import save_uploaded_file
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
from utils.file_reader import search_in_file

main_bp = Blueprint('main', 'main')

@main_bp.route('/')
def index():
    """Trang chủ"""
    if current_user.is_authenticated:
        return redirect(url_for('main.dashboard'))
    return redirect(url_for('auth.login'))

@main_bp.route('/dashboard')
@login_required
def dashboard():
    """Trang dashboard"""
    # Thống kê - Đếm tổng số trong database
    total_detai = DeTai.query.count()
    total_tailieu = TaiLieu.query.count()
    total_tiemluc = TiemLuc.query.count()
    total_sangkien = SangKien.query.count()
    total_users = User.query.count()
    
    # Đề tài gần đây
    recent_detai = DeTai.query.order_by(DeTai.created_at.desc()).limit(5).all()
    
    # Sáng kiến gần đây
    recent_sangkien = SangKien.query.order_by(SangKien.created_at.desc()).limit(5).all()
    
    return render_template('dashboard.html',
                         total_detai=total_detai,
                         total_giaotrinh=total_tailieu,  # Giữ tên biến cũ để tương thích template
                         total_tiemluc=total_tiemluc,
                         total_sangkien=total_sangkien,
                         total_users=total_users,
                         recent_detai=recent_detai,
                         recent_sangkien=recent_sangkien)

# ==================== QUẢN LÝ ĐỀ TÀI ====================
@main_bp.route('/detai')
@login_required
def detai_list():
    """Danh sách đề tài với phân trang và thống kê theo đơn vị"""
    page = request.args.get('page', 1, type=int)
    per_page = 20
    search = request.args.get('search', '')
    
    query = DeTai.query
    
    # Lọc theo đơn vị nếu không phải admin
    if current_user.role != 'admin':
        query = query.filter_by(don_vi=current_user.don_vi)
    
    # Tìm kiếm
    if search:
        query = query.filter(
            db.or_(
                DeTai.ma_de_tai.contains(search),
                DeTai.ten_de_tai.contains(search),
                DeTai.chu_nhiem.contains(search),
                DeTai.don_vi.contains(search)
            )
        )
    
    pagination = query.order_by(DeTai.created_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    
    # ========== THỐNG KÊ THEO ĐƠN VỊ ==========
    stats_query = db.session.query(
        DeTai.don_vi,
        func.count(DeTai.id).label('total')
    )
    
    if current_user.role != 'admin':
        stats_query = stats_query.filter(DeTai.don_vi == current_user.don_vi)
    
    if search:
        stats_query = stats_query.filter(
            db.or_(
                DeTai.ma_de_tai.contains(search),
                DeTai.ten_de_tai.contains(search),
                DeTai.chu_nhiem.contains(search),
                DeTai.don_vi.contains(search)
            )
        )
    
    donvi_stats = stats_query.group_by(DeTai.don_vi).order_by(func.count(DeTai.id).desc()).all()
    donvi_stats = [{'don_vi': dv or 'Chưa xác định', 'total': total} for dv, total in donvi_stats]
    
    # Thống kê theo cấp độ
    capdo_stats = db.session.query(
        DeTai.cap_do,
        func.count(DeTai.id).label('total')
    )
    if current_user.role != 'admin':
        capdo_stats = capdo_stats.filter(DeTai.don_vi == current_user.don_vi)
    if search:
        capdo_stats = capdo_stats.filter(
            db.or_(
                DeTai.ma_de_tai.contains(search),
                DeTai.ten_de_tai.contains(search),
                DeTai.chu_nhiem.contains(search)
            )
        )
    capdo_stats = capdo_stats.group_by(DeTai.cap_do).all()
    capdo_stats = [{'cap_do': cd or 'Chưa xác định', 'total': total} for cd, total in capdo_stats]
    
    # Thống kê theo trạng thái
    trangthai_stats = db.session.query(
        DeTai.trang_thai,
        func.count(DeTai.id).label('total')
    )
    if current_user.role != 'admin':
        trangthai_stats = trangthai_stats.filter(DeTai.don_vi == current_user.don_vi)
    if search:
        trangthai_stats = trangthai_stats.filter(
            db.or_(
                DeTai.ma_de_tai.contains(search),
                DeTai.ten_de_tai.contains(search),
                DeTai.chu_nhiem.contains(search)
            )
        )
    trangthai_stats = trangthai_stats.group_by(DeTai.trang_thai).all()
    trangthai_stats = [{'trang_thai': tt or 'Chưa xác định', 'total': total} for tt, total in trangthai_stats]
    
    # Tổng số
    total_detai = sum(s['total'] for s in donvi_stats)
    
    return render_template('detai/list.html', 
                         pagination=pagination, 
                         search=search,
                         donvi_stats=donvi_stats,
                         capdo_stats=capdo_stats,
                         trangthai_stats=trangthai_stats,
                         total_detai=total_detai)

@main_bp.route('/detai/add', methods=['GET', 'POST'])
@login_required
def detai_add():
    """Thêm đề tài mới - Tự động gán đơn vị của user"""
    if request.method == 'POST':
        # Chỉ cho phép thêm với đơn vị của mình (nếu không phải admin)
        don_vi = current_user.don_vi if current_user.role != 'admin' else request.form.get('don_vi')
        
        detai = DeTai(
            ma_de_tai=request.form.get('ma_de_tai'),
            ten_de_tai=request.form.get('ten_de_tai'),
            mo_ta=request.form.get('mo_ta'),
            linh_vuc=request.form.get('linh_vuc'),
            chu_nhiem=request.form.get('chu_nhiem'),
            don_vi=don_vi,
            trang_thai=request.form.get('trang_thai'),
            ket_qua = request.form.get('ket_qua'),
            created_by=current_user.id,
            cap_do=request.form.get('cap_do')
        )
      #  detai.ket_qua = request.form.get('ket_qua')
        # Xử lý ngày tháng
        ngay_bat_dau = request.form.get('ngay_bat_dau')
        ngay_ket_thuc = request.form.get('ngay_ket_thuc')
        if ngay_bat_dau:
            detai.ngay_bat_dau = datetime.strptime(ngay_bat_dau, '%Y-%m-%d').date()
        if ngay_ket_thuc:
            detai.ngay_ket_thuc = datetime.strptime(ngay_ket_thuc, '%Y-%m-%d').date()
        
        db.session.add(detai)
        db.session.flush()  # Để lấy ID
        
        # Tạo thư mục riêng cho đề tài: detai_<id>_<ma_de_tai>
        folder_name = f"detai_{detai.id}_{detai.ma_de_tai}"
          # Xử lý upload các file
        # Xử lý upload các file
        if 'file_sau_sua' in request.files:
            file = request.files['file_sau_sua']
            if file and file.filename:
                detai.file_sau_sua = save_uploaded_file(file, folder_name,category='detai')
        if 'bien_ban_ban' in request.files:
            file = request.files['bien_ban_ban']
            if file and file.filename:
                detai.bien_ban_ban = save_uploaded_file(file, folder_name,category='bien_ban_ban')
        
        if 'bien_ban_de_cuong' in request.files:
            file = request.files['bien_ban_de_cuong']
            if file and file.filename:
                detai.bien_ban_de_cuong = save_uploaded_file(file, folder_name,category='bien_ban_de_cuong')
        
        if 'bien_ban_nghiem_thu' in request.files:
            file = request.files['bien_ban_nghiem_thu']
            if file and file.filename:
                detai.bien_ban_nghiem_thu = save_uploaded_file(file, folder_name,category='bien_ban_nghiem_thu')
        
        # Xử lý thêm thành viên Ban đề tài
        ban_de_tai_text = request.form.get('ban_de_tai', '').strip()
        if ban_de_tai_text:
            for line in ban_de_tai_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    ban_member = BanDeTai(
                        de_tai_id=detai.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(ban_member)
        
        # Xử lý thêm thành viên Hội đồng duyệt đề cương
        hoi_dong_de_cuong_text = request.form.get('hoi_dong_de_cuong', '').strip()
        if hoi_dong_de_cuong_text:
            for line in hoi_dong_de_cuong_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    hd_member = HoiDongDeCuongDeTai(
                        de_tai_id=detai.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(hd_member)
        
        # Xử lý thêm thành viên Hội đồng nghiệm thu
        hoi_dong_nghiem_thu_text = request.form.get('hoi_dong_nghiem_thu', '').strip()
        if hoi_dong_nghiem_thu_text:
            for line in hoi_dong_nghiem_thu_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    hd_member = HoiDongNghiemThuDeTai(
                        de_tai_id=detai.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(hd_member)
        
        db.session.commit()
        
        flash('Thêm đề tài thành công!', 'success')
        return redirect(url_for('main.detai_list'))
    
    return render_template('detai/add.html', don_vi_list=DON_VI_LIST, cap_do_list=CAP_DO_LIST)

@main_bp.route('/detai/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def detai_edit(id):
    """Sửa đề tài - Chỉ sửa được đề tài của đơn vị mình hoặc admin"""
    detai = DeTai.query.get_or_404(id)
    
    # Kiểm tra quyền: chỉ admin hoặc user cùng đơn vị mới sửa được
    if current_user.role != 'admin' and detai.don_vi != current_user.don_vi:
        flash('Bạn không có quyền chỉnh sửa đề tài của đơn vị khác!', 'danger')
        return redirect(url_for('main.detai_list'))
    
    if request.method == 'POST':
        detai.ma_de_tai = request.form.get('ma_de_tai')
        detai.ten_de_tai = request.form.get('ten_de_tai')
        detai.mo_ta = request.form.get('mo_ta')
        detai.linh_vuc = request.form.get('linh_vuc')
        detai.chu_nhiem = request.form.get('chu_nhiem')
        detai.cap_do = request.form.get('cap_do')
        
        # Admin có thể đổi đơn vị
        if current_user.role == 'admin':
            detai.don_vi = request.form.get('don_vi')
        
        detai.trang_thai = request.form.get('trang_thai')
        detai.ket_qua = request.form.get('ket_qua')
        
        # Xử lý ngày tháng
        ngay_bat_dau = request.form.get('ngay_bat_dau')
        ngay_ket_thuc = request.form.get('ngay_ket_thuc')
        if ngay_bat_dau:
            detai.ngay_bat_dau = datetime.strptime(ngay_bat_dau, '%Y-%m-%d').date()
        if ngay_ket_thuc:
            detai.ngay_ket_thuc = datetime.strptime(ngay_ket_thuc, '%Y-%m-%d').date()
        
        # Tạo thư mục riêng cho đề tài
        folder_name = f"detai_{detai.id}_{detai.ma_de_tai}"
        
        # Xử lý upload các file (chỉ upload nếu có file mới)
        if 'bien_ban_ban' in request.files:
            file = request.files['bien_ban_ban']
            if file and file.filename:
                detai.bien_ban_ban = save_uploaded_file(file, folder_name,category='bien_ban_ban')
        
        if 'bien_ban_de_cuong' in request.files:
            file = request.files['bien_ban_de_cuong']
            if file and file.filename:
                detai.bien_ban_de_cuong = save_uploaded_file(file, folder_name,category='bien_ban_de_cuong')
        
        if 'bien_ban_nghiem_thu' in request.files:
            file = request.files['bien_ban_nghiem_thu']
            if file and file.filename:
                detai.bien_ban_nghiem_thu = save_uploaded_file(file, folder_name,category='bien_ban_nghiem_thu')
        
        if 'file_sau_sua' in request.files:
            file = request.files['file_sau_sua']
            if file and file.filename:
                detai.file_sau_sua = save_uploaded_file(file, folder_name,category='detai')
        
        # ======== CẬP NHẬT THÀNH VIÊN BAN ĐỀ TÀI ========
        # Xóa toàn bộ thành viên cũ
        BanDeTai.query.filter_by(de_tai_id=detai.id).delete()
        
        # Thêm lại từ form
        ban_de_tai_text = request.form.get('ban_de_tai', '').strip()
        if ban_de_tai_text:
            for line in ban_de_tai_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    ban_member = BanDeTai(
                        de_tai_id=detai.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(ban_member)
        
        # ======== CẬP NHẬT HỘI ĐỒNG DUYỆT ĐỀ CƯƠNG ========
        # Xóa toàn bộ thành viên cũ
        HoiDongDeCuongDeTai.query.filter_by(de_tai_id=detai.id).delete()
        
        # Thêm lại từ form
        hoi_dong_de_cuong_text = request.form.get('hoi_dong_de_cuong', '').strip()
        if hoi_dong_de_cuong_text:
            for line in hoi_dong_de_cuong_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    hd_member = HoiDongDeCuongDeTai(
                        de_tai_id=detai.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(hd_member)
        
        # ======== CẬP NHẬT HỘI ĐỒNG NGHIỆM THU ========
        # Xóa toàn bộ thành viên cũ
        HoiDongNghiemThuDeTai.query.filter_by(de_tai_id=detai.id).delete()
        
        # Thêm lại từ form
        hoi_dong_nghiem_thu_text = request.form.get('hoi_dong_nghiem_thu', '').strip()
        if hoi_dong_nghiem_thu_text:
            for line in hoi_dong_nghiem_thu_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    hd_member = HoiDongNghiemThuDeTai(
                        de_tai_id=detai.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(hd_member)
        
        db.session.commit()
        
        flash('Cập nhật đề tài thành công!', 'success')
        return redirect(url_for('main.detai_list'))
    
    return render_template('detai/edit.html', detai=detai, don_vi_list=DON_VI_LIST, cap_do_list=CAP_DO_LIST)

@main_bp.route('/detai/delete/<int:id>')
@login_required
def detai_delete(id):
    """Xóa đề tài - Chỉ xóa được đề tài của đơn vị mình hoặc admin"""
    detai = DeTai.query.get_or_404(id)
    
    # Kiểm tra quyền
    if current_user.role != 'admin' and detai.don_vi != current_user.don_vi:
        flash('Bạn không có quyền xóa đề tài của đơn vị khác!', 'danger')
        return redirect(url_for('main.detai_list'))
    
    db.session.delete(detai)
    db.session.commit()
    
    flash('Xóa đề tài thành công!', 'success')
    return redirect(url_for('main.detai_list'))

@main_bp.route('/tailieu/detail/<int:id>')
@login_required
def tailieu_detail(id):
    tailieu = TaiLieu.query.get_or_404(id)
    
    # Lấy danh sách thành viên ban và hội đồng
    ban_members = BanTaiLieu.query.filter_by(tai_lieu_id=id).all()
    hoi_dong_de_cuong_members = HoiDongDeCuongTaiLieu.query.filter_by(tai_lieu_id=id).all()
    hoi_dong_nghiem_thu_members = HoiDongNghiemThuTaiLieu.query.filter_by(tai_lieu_id=id).all()
    
    return render_template('tailieu/detail.html', 
                         tailieu=tailieu,
                         ban_members=ban_members,
                         hoi_dong_de_cuong_members=hoi_dong_de_cuong_members,
                         hoi_dong_nghiem_thu_members=hoi_dong_nghiem_thu_members)
# ==================== QUẢN LÝ TÀI LIỆU ====================
@main_bp.route('/tailieu')
@login_required
def tailieu_list():
    """Danh sách tài liệu với phân trang và thống kê theo đơn vị"""
    page = request.args.get('page', 1, type=int)
    per_page = 20
    search = request.args.get('search', '')
    
    query = TaiLieu.query
    
    # Lọc theo đơn vị nếu không phải admin
    if current_user.role != 'admin':
        query = query.filter_by(don_vi=current_user.don_vi)
    
    # Tìm kiếm
    if search:
        query = query.filter(
            db.or_(
                TaiLieu.ma_tai_lieu.contains(search),
                TaiLieu.ten_tai_lieu.contains(search),
                TaiLieu.tac_gia.contains(search),
                TaiLieu.don_vi.contains(search)
            )
        )
    
    pagination = query.order_by(TaiLieu.created_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    
    # ========== THỐNG KÊ THEO ĐƠN VỊ ==========
    stats_query = db.session.query(
        TaiLieu.don_vi,
        func.count(TaiLieu.id).label('total')
    )
    
    # Nếu không phải admin, chỉ thống kê đơn vị của mình
    if current_user.role != 'admin':
        stats_query = stats_query.filter(TaiLieu.don_vi == current_user.don_vi)
    
    # Nếu có tìm kiếm, áp dụng vào thống kê
    if search:
        stats_query = stats_query.filter(
            db.or_(
                TaiLieu.ma_tai_lieu.contains(search),
                TaiLieu.ten_tai_lieu.contains(search),
                TaiLieu.tac_gia.contains(search),
                TaiLieu.don_vi.contains(search)
            )
        )
    
    donvi_stats = stats_query.group_by(TaiLieu.don_vi).order_by(func.count(TaiLieu.id).desc()).all()
    donvi_stats = [{'don_vi': dv or 'Chưa xác định', 'total': total} for dv, total in donvi_stats]
    
    # Thống kê theo loại
    loai_stats = db.session.query(
        TaiLieu.loai,
        func.count(TaiLieu.id).label('total')
    )
    if current_user.role != 'admin':
        loai_stats = loai_stats.filter(TaiLieu.don_vi == current_user.don_vi)
    if search:
        loai_stats = loai_stats.filter(
            db.or_(
                TaiLieu.ma_tai_lieu.contains(search),
                TaiLieu.ten_tai_lieu.contains(search),
                TaiLieu.tac_gia.contains(search)
            )
        )
    loai_stats = loai_stats.group_by(TaiLieu.loai).all()
    loai_stats = [{'loai': loai or 'Chưa xác định', 'total': total} for loai, total in loai_stats]
    
    # Tổng số
    total_tailieu = sum(s['total'] for s in donvi_stats)
    
    return render_template('tailieu/list.html', 
                         pagination=pagination, 
                         search=search,
                         donvi_stats=donvi_stats,
                         loai_stats=loai_stats,
                         total_tailieu=total_tailieu)
# Route tương thích với tên cũ
@main_bp.route('/giaotrinh')
@login_required
def giaotrinh_list():
    """Chuyển hướng sang tài liệu"""
    return redirect(url_for('main.tailieu_list'))

@main_bp.route('/tailieu/add', methods=['GET', 'POST'])
@login_required
def tailieu_add():
    if request.method == 'POST':
        # Tạo mã tài liệu tự động nếu không nhập
        ma_tai_lieu = request.form.get('ma_tai_lieu')
        if not ma_tai_lieu:
            year = datetime.now().year
            count = TaiLieu.query.filter(TaiLieu.ma_tai_lieu.like(f'TL{year}%')).count()
            ma_tai_lieu = f'TL{year}{count+1:04d}'
        
        tailieu = TaiLieu(
            ma_tai_lieu=ma_tai_lieu,
            ten_tai_lieu=request.form.get('ten_tai_lieu'),
            loai=request.form.get('loai'),
            tac_gia=request.form.get('tac_gia', '').strip(),
            mon_hoc=request.form.get('mon_hoc'),
            linh_vuc=request.form.get('linh_vuc'),
            nha_xuat_ban=request.form.get('nha_xuat_ban'),
            nam_xuat_ban=request.form.get('nam_xuat_ban') or None,
            trang_thai=request.form.get('trang_thai', 'Đang thực hiện'),
            ket_qua=request.form.get('ket_qua'),
            mo_ta=request.form.get('mo_ta'),
            don_vi=current_user.don_vi
        )
        
        # Ngày tháng
        ngay_bat_dau = request.form.get('ngay_bat_dau')
        if ngay_bat_dau:
            tailieu.ngay_bat_dau = datetime.strptime(ngay_bat_dau, '%Y-%m-%d').date()
        
        ngay_ket_thuc = request.form.get('ngay_ket_thuc')
        if ngay_ket_thuc:
            tailieu.ngay_ket_thuc = datetime.strptime(ngay_ket_thuc, '%Y-%m-%d').date()
        
        # Thêm vào DB trước để có ID
        db.session.add(tailieu)
        db.session.flush()  # Lấy ID nhưng chưa commit
        
        # ======== THÊM BAN TÀI LIỆU ========
        ban_tai_lieu_text = request.form.get('ban_tai_lieu', '').strip()
        if ban_tai_lieu_text:
            for line in ban_tai_lieu_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip() if len(parts) > 1 else ''
                    if ho_ten:
                        ban_member = BanTaiLieu(
                            tai_lieu_id=tailieu.id,
                            ho_ten=ho_ten,
                            vai_tro=vai_tro
                        )
                        db.session.add(ban_member)
        
        # ======== THÊM HỘI ĐỒNG DUYỆT ĐỀ CƯƠNG ========
        hoi_dong_de_cuong_text = request.form.get('hoi_dong_de_cuong', '').strip()
        if hoi_dong_de_cuong_text:
            for line in hoi_dong_de_cuong_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip() if len(parts) > 1 else ''
                    if ho_ten:
                        member = HoiDongDeCuongTaiLieu(
                            tai_lieu_id=tailieu.id,
                            ho_ten=ho_ten,
                            vai_tro=vai_tro
                        )
                        db.session.add(member)
        
        # ======== THÊM HỘI ĐỒNG NGHIỆM THU ========
        hoi_dong_nghiem_thu_text = request.form.get('hoi_dong_nghiem_thu', '').strip()
        if hoi_dong_nghiem_thu_text:
            for line in hoi_dong_nghiem_thu_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip() if len(parts) > 1 else ''
                    if ho_ten:
                        member = HoiDongNghiemThuTaiLieu(
                            tai_lieu_id=tailieu.id,
                            ho_ten=ho_ten,
                            vai_tro=vai_tro
                        )
                        db.session.add(member)
        
        # Upload files
        folder_name = f"tailieu_{ma_tai_lieu}"
        
        if 'file_tai_lieu' in request.files:
            file = request.files['file_tai_lieu']
            if file and file.filename:
                tailieu.file_tai_lieu = save_uploaded_file(file, folder_name, category='tailieu')
        
        if 'bien_ban_ban' in request.files:
            file = request.files['bien_ban_ban']
            if file and file.filename:
                tailieu.bien_ban_ban = save_uploaded_file(file, folder_name, category='bien_ban_ban')
        
        if 'bien_ban_de_cuong' in request.files:
            file = request.files['bien_ban_de_cuong']
            if file and file.filename:
                tailieu.bien_ban_de_cuong = save_uploaded_file(file, folder_name, category='bien_ban_de_cuong')
        
        if 'bien_ban_nghiem_thu' in request.files:
            file = request.files['bien_ban_nghiem_thu']
            if file and file.filename:
                tailieu.bien_ban_nghiem_thu = save_uploaded_file(file, folder_name, category='bien_ban_nghiem_thu')
        
        try:
            db.session.commit()
            flash('Thêm tài liệu thành công!', 'success')
            return redirect(url_for('main.tailieu_list'))
        except Exception as e:
            db.session.rollback()
            flash(f'Lỗi khi thêm tài liệu: {str(e)}', 'danger')
    
    return render_template('tailieu/add.html',don_vi_list=DON_VI_LIST, cap_do_list=CAP_DO_LIST)

# Route tương thích
@main_bp.route('/giaotrinh/add', methods=['GET', 'POST'])
@login_required
def giaotrinh_add():
    """Chuyển hướng sang thêm tài liệu"""
    if request.method == 'POST':
        return tailieu_add()
    return redirect(url_for('main.tailieu_add'))

@main_bp.route('/tailieu/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def tailieu_edit(id):
    tailieu = TaiLieu.query.get_or_404(id)
    
    # Kiểm tra quyền
    if current_user.role != 'admin' and current_user.don_vi != tailieu.don_vi:
        flash('Bạn không có quyền chỉnh sửa tài liệu này!', 'danger')
        return redirect(url_for('main.tailieu_list'))
    
    if request.method == 'POST':
        # Cập nhật thông tin cơ bản
        tailieu.ma_tai_lieu = request.form.get('ma_tai_lieu')
        tailieu.ten_tai_lieu = request.form.get('ten_tai_lieu')
        tailieu.loai = request.form.get('loai')
        tailieu.tac_gia = request.form.get('tac_gia', '').strip()
        tailieu.mon_hoc = request.form.get('mon_hoc')
        tailieu.linh_vuc = request.form.get('linh_vuc')
        tailieu.nha_xuat_ban = request.form.get('nha_xuat_ban')
        tailieu.nam_xuat_ban = request.form.get('nam_xuat_ban') or None
        tailieu.trang_thai = request.form.get('trang_thai')
        tailieu.ket_qua = request.form.get('ket_qua')
        tailieu.mo_ta = request.form.get('mo_ta')
        
        # Ngày tháng
        ngay_bat_dau = request.form.get('ngay_bat_dau')
        if ngay_bat_dau:
            tailieu.ngay_bat_dau = datetime.strptime(ngay_bat_dau, '%Y-%m-%d').date()
        
        ngay_ket_thuc = request.form.get('ngay_ket_thuc')
        if ngay_ket_thuc:
            tailieu.ngay_ket_thuc = datetime.strptime(ngay_ket_thuc, '%Y-%m-%d').date()
        
        # ======== CẬP NHẬT BAN TÀI LIỆU ========
        # Xóa toàn bộ thành viên cũ
        BanTaiLieu.query.filter_by(tai_lieu_id=tailieu.id).delete()
        
        # Thêm lại từ form
        ban_tai_lieu_text = request.form.get('ban_tai_lieu', '').strip()
        if ban_tai_lieu_text:
            for line in ban_tai_lieu_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip() if len(parts) > 1 else ''
                    if ho_ten:
                        ban_member = BanTaiLieu(
                            tai_lieu_id=tailieu.id,
                            ho_ten=ho_ten,
                            vai_tro=vai_tro
                        )
                        db.session.add(ban_member)
        
        # ======== CẬP NHẬT HỘI ĐỒNG DUYỆT ĐỀ CƯƠNG ========
        HoiDongDeCuongTaiLieu.query.filter_by(tai_lieu_id=tailieu.id).delete()
        
        hoi_dong_de_cuong_text = request.form.get('hoi_dong_de_cuong', '').strip()
        if hoi_dong_de_cuong_text:
            for line in hoi_dong_de_cuong_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip() if len(parts) > 1 else ''
                    if ho_ten:
                        member = HoiDongDeCuongTaiLieu(
                            tai_lieu_id=tailieu.id,
                            ho_ten=ho_ten,
                            vai_tro=vai_tro
                        )
                        db.session.add(member)
        
        # ======== CẬP NHẬT HỘI ĐỒNG NGHIỆM THU ========
        HoiDongNghiemThuTaiLieu.query.filter_by(tai_lieu_id=tailieu.id).delete()
        
        hoi_dong_nghiem_thu_text = request.form.get('hoi_dong_nghiem_thu', '').strip()
        if hoi_dong_nghiem_thu_text:
            for line in hoi_dong_nghiem_thu_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip() if len(parts) > 1 else ''
                    if ho_ten:
                        member = HoiDongNghiemThuTaiLieu(
                            tai_lieu_id=tailieu.id,
                            ho_ten=ho_ten,
                            vai_tro=vai_tro
                        )
                        db.session.add(member)
        
        # ======== XỬ LÝ UPLOAD FILE ========
        folder_name = f"tailieu_{tailieu.ma_tai_lieu}"
        
        # File tài liệu chính
        if 'file_tai_lieu' in request.files:
            file = request.files['file_tai_lieu']
            if file and file.filename:
                tailieu.file_tai_lieu = save_uploaded_file(file, folder_name, category='tailieu')
        
        # Các biên bản
        if 'bien_ban_ban' in request.files:
            file = request.files['bien_ban_ban']
            if file and file.filename:
                tailieu.bien_ban_ban = save_uploaded_file(file, folder_name, category='bien_ban_ban')
        
        if 'bien_ban_de_cuong' in request.files:
            file = request.files['bien_ban_de_cuong']
            if file and file.filename:
                tailieu.bien_ban_de_cuong = save_uploaded_file(file, folder_name, category='bien_ban_de_cuong')
        
        if 'bien_ban_nghiem_thu' in request.files:
            file = request.files['bien_ban_nghiem_thu']
            if file and file.filename:
                tailieu.bien_ban_nghiem_thu = save_uploaded_file(file, folder_name, category='bien_ban_nghiem_thu')
        
        try:
            db.session.commit()
            flash('Cập nhật tài liệu thành công!', 'success')
            return redirect(url_for('main.tailieu_detail', id=tailieu.id))
        except Exception as e:
            db.session.rollback()
            flash(f'Lỗi khi cập nhật: {str(e)}', 'danger')
    
    return render_template('tailieu/edit.html', tailieu=tailieu)


# Route tương thích
@main_bp.route('/giaotrinh/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def giaotrinh_edit(id):
    """Chuyển hướng sang sửa tài liệu"""
    if request.method == 'POST':
        return tailieu_edit(id)
    return redirect(url_for('main.tailieu_edit', id=id))

@main_bp.route('/tailieu/delete/<int:id>')
@login_required
def tailieu_delete(id):
    """Xóa tài liệu"""
    tailieu = TaiLieu.query.get_or_404(id)
    
    if current_user.role != 'admin' and tailieu.don_vi != current_user.don_vi:
        flash('Bạn không có quyền xóa tài liệu của đơn vị khác!', 'danger')
        return redirect(url_for('main.tailieu_list'))
    
    db.session.delete(tailieu)
    db.session.commit()
    
    flash('Xóa tài liệu thành công!', 'success')
    return redirect(url_for('main.tailieu_list'))

# Route tương thích
@main_bp.route('/giaotrinh/delete/<int:id>')
@login_required
def giaotrinh_delete(id):
    """Chuyển hướng sang xóa tài liệu"""
    return tailieu_delete(id)

# ==================== QUẢN LÝ TIỀM LỰC ====================
@main_bp.route('/tiemluc')
@login_required
def tiemluc_list():
    """Danh sách tiềm lực với phân trang và thống kê theo đơn vị"""
    page = request.args.get('page', 1, type=int)
    per_page = 20
    search = request.args.get('search', '')
    
    query = TiemLuc.query
    
    # Lọc theo đơn vị nếu không phải admin
    if current_user.role != 'admin':
        query = query.filter_by(don_vi=current_user.don_vi)
    
    # Tìm kiếm
    if search:
        query = query.filter(
            db.or_(
                TiemLuc.ho_ten.contains(search),
                TiemLuc.hoc_vi.contains(search),
                TiemLuc.hoc_ham.contains(search),
                TiemLuc.don_vi.contains(search)
            )
        )
    
    pagination = query.order_by(TiemLuc.created_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    
    # ========== THỐNG KÊ THEO ĐƠN VỊ ==========
    stats_query = db.session.query(
        TiemLuc.don_vi,
        func.count(TiemLuc.id).label('total')
    )
    
    if current_user.role != 'admin':
        stats_query = stats_query.filter(TiemLuc.don_vi == current_user.don_vi)
    
    if search:
        stats_query = stats_query.filter(
            db.or_(
                TiemLuc.ho_ten.contains(search),
                TiemLuc.hoc_vi.contains(search),
                TiemLuc.hoc_ham.contains(search),
                TiemLuc.don_vi.contains(search)
            )
        )
    
    donvi_stats = stats_query.group_by(TiemLuc.don_vi).order_by(func.count(TiemLuc.id).desc()).all()
    donvi_stats = [{'don_vi': dv or 'Chưa xác định', 'total': total} for dv, total in donvi_stats]
    
    # Thống kê theo học vị
    hocvi_stats = db.session.query(
        TiemLuc.hoc_vi,
        func.count(TiemLuc.id).label('total')
    ).filter(TiemLuc.hoc_vi.isnot(None))
    if current_user.role != 'admin':
        hocvi_stats = hocvi_stats.filter(TiemLuc.don_vi == current_user.don_vi)
    if search:
        hocvi_stats = hocvi_stats.filter(
            db.or_(
                TiemLuc.ho_ten.contains(search),
                TiemLuc.don_vi.contains(search)
            )
        )
    hocvi_stats = hocvi_stats.group_by(TiemLuc.hoc_vi).all()
    hocvi_stats = [{'hoc_vi': hv or 'Chưa có', 'total': total} for hv, total in hocvi_stats]
    
    # Thống kê theo học hàm
    hocham_stats = db.session.query(
        TiemLuc.hoc_ham,
        func.count(TiemLuc.id).label('total')
    ).filter(TiemLuc.hoc_ham.isnot(None))
    if current_user.role != 'admin':
        hocham_stats = hocham_stats.filter(TiemLuc.don_vi == current_user.don_vi)
    if search:
        hocham_stats = hocham_stats.filter(
            db.or_(
                TiemLuc.ho_ten.contains(search),
                TiemLuc.don_vi.contains(search)
            )
        )
    hocham_stats = hocham_stats.group_by(TiemLuc.hoc_ham).all()
    hocham_stats = [{'hoc_ham': hh or 'Chưa có', 'total': total} for hh, total in hocham_stats]
    
    # Tổng số
    total_tiemluc = sum(s['total'] for s in donvi_stats)
    
    return render_template('tiemluc/list.html', 
                         pagination=pagination, 
                         search=search,
                         donvi_stats=donvi_stats,
                         hocvi_stats=hocvi_stats,
                         hocham_stats=hocham_stats,
                         total_tiemluc=total_tiemluc)

@main_bp.route('/tiemluc/add', methods=['GET', 'POST'])
@login_required
def tiemluc_add():
    """Thêm tiềm lực mới - Tự động gán đơn vị của user"""
    if request.method == 'POST':
        # Chỉ cho phép thêm với đơn vị của mình (nếu không phải admin)
        don_vi = current_user.don_vi if current_user.role != 'admin' else request.form.get('don_vi')
        
        tiemluc = TiemLuc(
            ma_can_bo=request.form.get('ma_can_bo'),
            ho_ten=request.form.get('ho_ten'),
            gioi_tinh=request.form.get('gioi_tinh'),
            don_vi=don_vi,
            chuc_vu=request.form.get('chuc_vu'),
            hoc_vi=request.form.get('hoc_vi'),
            hoc_ham=request.form.get('hoc_ham'),
            chuyen_mon=request.form.get('chuyen_mon'),
            email=request.form.get('email'),
            dien_thoai=request.form.get('dien_thoai'),
            kinh_nghiem=request.form.get('kinh_nghiem'),
            thanh_tich=request.form.get('thanh_tich'),
            created_by=current_user.id
        )
        
        ngay_sinh = request.form.get('ngay_sinh')
        if ngay_sinh:
            tiemluc.ngay_sinh = datetime.strptime(ngay_sinh, '%Y-%m-%d').date()
        
        db.session.add(tiemluc)
        db.session.commit()
        
        flash('Thêm tiềm lực khoa học thành công!', 'success')
        return redirect(url_for('main.tiemluc_list'))
    
    return render_template('tiemluc/add.html', don_vi_list=DON_VI_LIST)

@main_bp.route('/tiemluc/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def tiemluc_edit(id):
    """Sửa tiềm lực - Chỉ sửa được cán bộ của đơn vị mình hoặc admin"""
    tiemluc = TiemLuc.query.get_or_404(id)
    
    # Kiểm tra quyền: chỉ admin hoặc user cùng đơn vị mới sửa được
    if current_user.role != 'admin' and tiemluc.don_vi != current_user.don_vi:
        flash('Bạn không có quyền chỉnh sửa cán bộ của đơn vị khác!', 'danger')
        return redirect(url_for('main.tiemluc_list'))
    
    if request.method == 'POST':
        tiemluc.ma_can_bo = request.form.get('ma_can_bo')
        tiemluc.ho_ten = request.form.get('ho_ten')
        tiemluc.gioi_tinh = request.form.get('gioi_tinh')
        
        # Admin có thể đổi đơn vị
        if current_user.role == 'admin':
            tiemluc.don_vi = request.form.get('don_vi')
        
        tiemluc.chuc_vu = request.form.get('chuc_vu')
        tiemluc.hoc_vi = request.form.get('hoc_vi')
        tiemluc.hoc_ham = request.form.get('hoc_ham')
        tiemluc.chuyen_mon = request.form.get('chuyen_mon')
        tiemluc.email = request.form.get('email')
        tiemluc.dien_thoai = request.form.get('dien_thoai')
        tiemluc.kinh_nghiem = request.form.get('kinh_nghiem')
        tiemluc.thanh_tich = request.form.get('thanh_tich')
        
        ngay_sinh = request.form.get('ngay_sinh')
        if ngay_sinh:
            tiemluc.ngay_sinh = datetime.strptime(ngay_sinh, '%Y-%m-%d').date()
        
        db.session.commit()
        
        flash('Cập nhật tiềm lực khoa học thành công!', 'success')
        return redirect(url_for('main.tiemluc_list'))
    
    return render_template('tiemluc/edit.html', tiemluc=tiemluc, don_vi_list=DON_VI_LIST)

@main_bp.route('/tiemluc/delete/<int:id>')
@login_required
def tiemluc_delete(id):
    """Xóa tiềm lực - Chỉ xóa được cán bộ của đơn vị mình hoặc admin"""
    tiemluc = TiemLuc.query.get_or_404(id)
    
    # Kiểm tra quyền
    if current_user.role != 'admin' and tiemluc.don_vi != current_user.don_vi:
        flash('Bạn không có quyền xóa cán bộ của đơn vị khác!', 'danger')
        return redirect(url_for('main.tiemluc_list'))
    
    db.session.delete(tiemluc)
    db.session.commit()
    
    flash('Xóa tiềm lực khoa học thành công!', 'success')
    return redirect(url_for('main.tiemluc_list'))

# ==================== QUẢN LÝ SÁNG KIẾN ====================
@main_bp.route('/sangkien')
@login_required
def sangkien_list():
    """Danh sách sáng kiến với phân trang và thống kê theo đơn vị"""
    page = request.args.get('page', 1, type=int)
    per_page = 20
    search = request.args.get('search', '')
    
    query = SangKien.query
    
    # Lọc theo đơn vị nếu không phải admin
    if current_user.role != 'admin':
        query = query.filter_by(don_vi=current_user.don_vi)
    
    # Tìm kiếm
    if search:
        query = query.filter(
            db.or_(
                SangKien.ma_sang_kien.contains(search),
                SangKien.ten_sang_kien.contains(search),
                SangKien.chu_nhiem.contains(search),
                SangKien.don_vi.contains(search)
            )
        )
    
    pagination = query.order_by(SangKien.created_at.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    
    # ========== THỐNG KÊ THEO ĐƠN VỊ ==========
    stats_query = db.session.query(
        SangKien.don_vi,
        func.count(SangKien.id).label('total')
    )
    
    if current_user.role != 'admin':
        stats_query = stats_query.filter(SangKien.don_vi == current_user.don_vi)
    
    if search:
        stats_query = stats_query.filter(
            db.or_(
                SangKien.ma_sang_kien.contains(search),
                SangKien.ten_sang_kien.contains(search),
                SangKien.chu_nhiem.contains(search),
                SangKien.don_vi.contains(search)
            )
        )
    
    donvi_stats = stats_query.group_by(SangKien.don_vi).order_by(func.count(SangKien.id).desc()).all()
    donvi_stats = [{'don_vi': dv or 'Chưa xác định', 'total': total} for dv, total in donvi_stats]
    
    # Thống kê theo trạng thái
    trangthai_stats = db.session.query(
        SangKien.trang_thai,
        func.count(SangKien.id).label('total')
    )
    if current_user.role != 'admin':
        trangthai_stats = trangthai_stats.filter(SangKien.don_vi == current_user.don_vi)
    if search:
        trangthai_stats = trangthai_stats.filter(
            db.or_(
                SangKien.ma_sang_kien.contains(search),
                SangKien.ten_sang_kien.contains(search),
                SangKien.chu_nhiem.contains(search)
            )
        )
    trangthai_stats = trangthai_stats.group_by(SangKien.trang_thai).all()
    trangthai_stats = [{'trang_thai': tt or 'Chưa xác định', 'total': total} for tt, total in trangthai_stats]
    
    # Tổng số
    total_sangkien = sum(s['total'] for s in donvi_stats)
    
    return render_template('sangkien/list.html', 
                         pagination=pagination, 
                         search=search,
                         donvi_stats=donvi_stats,
                         trangthai_stats=trangthai_stats,
                         total_sangkien=total_sangkien)

@main_bp.route('/sangkien/add', methods=['GET', 'POST'])
@login_required
def sangkien_add():
    """Thêm sáng kiến mới - Tự động gán đơn vị của user"""
    if request.method == 'POST':
        # Kiểm tra mã sáng kiến đã tồn tại chưa
        ma_sang_kien = request.form.get('ma_sang_kien')
        existing = SangKien.query.filter_by(ma_sang_kien=ma_sang_kien).first()
        if existing:
            flash(f'Mã sáng kiến "{ma_sang_kien}" đã tồn tại! Vui lòng sử dụng mã khác.', 'danger')
            return render_template('sangkien/add.html', don_vi_list=DON_VI_LIST, form_data=request.form)
        
        # Chỉ cho phép thêm với đơn vị của mình (nếu không phải admin)
        don_vi = current_user.don_vi if current_user.role != 'admin' else request.form.get('don_vi')
        
        sangkien = SangKien(
            ma_sang_kien=ma_sang_kien,
            ten_sang_kien=request.form.get('ten_sang_kien'),
            mo_ta=request.form.get('mo_ta'),
            linh_vuc=request.form.get('linh_vuc'),
            chu_nhiem=request.form.get('chu_nhiem'),
            don_vi=don_vi,
            cap_do=request.form.get('cap_do'),
            trang_thai=request.form.get('trang_thai', 'Đang thực hiện'),
            ket_qua=request.form.get('ket_qua'),
            created_by=current_user.id
        )
        
        # Xử lý ngày tháng
        ngay_bat_dau = request.form.get('ngay_bat_dau')
        ngay_ket_thuc = request.form.get('ngay_ket_thuc')
        sangkien.ket_qua = request.form.get('ket_qua')
        if ngay_bat_dau:
            sangkien.ngay_bat_dau = datetime.strptime(ngay_bat_dau, '%Y-%m-%d').date()
        if ngay_ket_thuc:
            sangkien.ngay_ket_thuc = datetime.strptime(ngay_ket_thuc, '%Y-%m-%d').date()
        
        db.session.add(sangkien)
        db.session.flush()  # Để lấy ID
        
        # Tạo thư mục riêng cho sáng kiến: sangkien_<id>_<ma_sang_kien>
        folder_name = f"sangkien_{sangkien.id}_{sangkien.ma_sang_kien}"
        
        # Xử lý upload các file
        if 'bien_ban_ban' in request.files:
            file = request.files['bien_ban_ban']
            if file and file.filename:
                sangkien.bien_ban_ban = save_uploaded_file(file, folder_name,category='bien_ban_ban')
        
        if 'bien_ban_de_cuong' in request.files:
            file = request.files['bien_ban_de_cuong']
            if file and file.filename:
                sangkien.bien_ban_de_cuong = save_uploaded_file(file, folder_name,category='bien_ban_de_cuong')
        
        if 'bien_ban_nghiem_thu' in request.files:
            file = request.files['bien_ban_nghiem_thu']
            if file and file.filename:
                sangkien.bien_ban_nghiem_thu = save_uploaded_file(file, folder_name,category='bien_ban_nghiem_thu')
        
        if 'file_sau_sua' in request.files:
            file = request.files['file_sau_sua']
            if file and file.filename:
                sangkien.file_sau_sua = save_uploaded_file(file, folder_name,category='sangkien')
        
        # Xử lý thêm thành viên Ban sáng kiến
        ban_sang_kien_text = request.form.get('ban_sang_kien', '').strip()
        if ban_sang_kien_text:
            for line in ban_sang_kien_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    ban_member = BanSangKien(
                        sang_kien_id=sangkien.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(ban_member)
        
        # Xử lý thêm thành viên Hội đồng duyệt đề cương
        hoi_dong_de_cuong_text = request.form.get('hoi_dong_de_cuong', '').strip()
        if hoi_dong_de_cuong_text:
            for line in hoi_dong_de_cuong_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    hd_member = HoiDongDeCuongSangKien(
                        sang_kien_id=sangkien.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(hd_member)
        
        # Xử lý thêm thành viên Hội đồng nghiệm thu
        hoi_dong_nghiem_thu_text = request.form.get('hoi_dong_nghiem_thu', '').strip()
        if hoi_dong_nghiem_thu_text:
            for line in hoi_dong_nghiem_thu_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    hd_member = HoiDongNghiemThuSangKien(
                        sang_kien_id=sangkien.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(hd_member)
        
        db.session.commit()
        
        flash('Thêm sáng kiến thành công!', 'success')
        return redirect(url_for('main.sangkien_list'))
    
    return render_template('sangkien/add.html', don_vi_list=DON_VI_LIST, cap_do_list=CAP_DO_LIST)

@main_bp.route('/sangkien/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def sangkien_edit(id):
    """Sửa sáng kiến - Chỉ sửa được sáng kiến của đơn vị mình hoặc admin"""
    sangkien = SangKien.query.get_or_404(id)
    
    # Kiểm tra quyền: chỉ admin hoặc user cùng đơn vị mới sửa được
    if current_user.role != 'admin' and sangkien.don_vi != current_user.don_vi:
        flash('Bạn không có quyền chỉnh sửa sáng kiến của đơn vị khác!', 'danger')
        return redirect(url_for('main.sangkien_list'))
    
    if request.method == 'POST':
        sangkien.ma_sang_kien = request.form.get('ma_sang_kien')
        sangkien.ten_sang_kien = request.form.get('ten_sang_kien')
        sangkien.mo_ta = request.form.get('mo_ta')
        sangkien.linh_vuc = request.form.get('linh_vuc')
        sangkien.chu_nhiem = request.form.get('chu_nhiem')
        sangkien.cap_do = request.form.get('cap_do')
        
        # Admin có thể đổi đơn vị
        if current_user.role == 'admin':
            sangkien.don_vi = request.form.get('don_vi')
        
        sangkien.trang_thai = request.form.get('trang_thai')
        sangkien.ket_qua = request.form.get('ket_qua')
        
        # Xử lý ngày tháng
        ngay_bat_dau = request.form.get('ngay_bat_dau')
        ngay_ket_thuc = request.form.get('ngay_ket_thuc')
        if ngay_bat_dau:
            sangkien.ngay_bat_dau = datetime.strptime(ngay_bat_dau, '%Y-%m-%d').date()
        if ngay_ket_thuc:
            sangkien.ngay_ket_thuc = datetime.strptime(ngay_ket_thuc, '%Y-%m-%d').date()
        
        # Tạo thư mục riêng cho sáng kiến
        folder_name = f"sangkien_{sangkien.id}_{sangkien.ma_sang_kien}"
        
        # Xử lý upload các file (chỉ upload nếu có file mới)
        if 'bien_ban_ban' in request.files:
            file = request.files['bien_ban_ban']
            if file and file.filename:
                sangkien.bien_ban_ban = save_uploaded_file(file, folder_name,category='bien_ban_ban')
        
        if 'bien_ban_de_cuong' in request.files:
            file = request.files['bien_ban_de_cuong']
            if file and file.filename:
                sangkien.bien_ban_de_cuong = save_uploaded_file(file, folder_name,category='bien_ban_de_cuong')
        
        if 'bien_ban_nghiem_thu' in request.files:
            file = request.files['bien_ban_nghiem_thu']
            if file and file.filename:
                sangkien.bien_ban_nghiem_thu = save_uploaded_file(file, folder_name,category='bien_ban_nghiem_thu')
        
        if 'file_sau_sua' in request.files:
            file = request.files['file_sau_sua']
            if file and file.filename:
                sangkien.file_sau_sua = save_uploaded_file(file, folder_name,category='sangkien')
        
        # ======== CẬP NHẬT THÀNH VIÊN BAN SÁNG KIẾN ========
        # Xóa toàn bộ thành viên cũ
        BanSangKien.query.filter_by(sang_kien_id=sangkien.id).delete()
        
        # Thêm lại từ form
        ban_sang_kien_text = request.form.get('ban_sang_kien', '').strip()
        if ban_sang_kien_text:
            for line in ban_sang_kien_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    ban_member = BanSangKien(
                        sang_kien_id=sangkien.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(ban_member)
        
        # ======== CẬP NHẬT HỘI ĐỒNG DUYỆT ĐỀ CƯƠNG ========
        # Xóa toàn bộ thành viên cũ
        HoiDongDeCuongSangKien.query.filter_by(sang_kien_id=sangkien.id).delete()
        
        # Thêm lại từ form
        hoi_dong_de_cuong_text = request.form.get('hoi_dong_de_cuong', '').strip()
        if hoi_dong_de_cuong_text:
            for line in hoi_dong_de_cuong_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    hd_member = HoiDongDeCuongSangKien(
                        sang_kien_id=sangkien.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(hd_member)
        
        # ======== CẬP NHẬT HỘI ĐỒNG NGHIỆM THU ========
        # Xóa toàn bộ thành viên cũ
        HoiDongNghiemThuSangKien.query.filter_by(sang_kien_id=sangkien.id).delete()
        
        # Thêm lại từ form
        hoi_dong_nghiem_thu_text = request.form.get('hoi_dong_nghiem_thu', '').strip()
        if hoi_dong_nghiem_thu_text:
            for line in hoi_dong_nghiem_thu_text.split('\n'):
                line = line.strip()
                if line and '-' in line:
                    parts = line.split('-', 1)
                    vai_tro = parts[0].strip()
                    ho_ten = parts[1].strip()
                    hd_member = HoiDongNghiemThuSangKien(
                        sang_kien_id=sangkien.id,
                        vai_tro=vai_tro,
                        ho_ten=ho_ten
                    )
                    db.session.add(hd_member)
        
        db.session.commit()
        
        flash('Cập nhật sáng kiến thành công!', 'success')
        return redirect(url_for('main.sangkien_list'))
    
    return render_template('sangkien/edit.html', sangkien=sangkien, don_vi_list=DON_VI_LIST, cap_do_list=CAP_DO_LIST)

@main_bp.route('/sangkien/delete/<int:id>')
@login_required
def sangkien_delete(id):
    """Xóa sáng kiến - Chỉ xóa được sáng kiến của đơn vị mình hoặc admin"""
    sangkien = SangKien.query.get_or_404(id)
    
    # Kiểm tra quyền
    if current_user.role != 'admin' and sangkien.don_vi != current_user.don_vi:
        flash('Bạn không có quyền xóa sáng kiến của đơn vị khác!', 'danger')
        return redirect(url_for('main.sangkien_list'))
    
    db.session.delete(sangkien)
    db.session.commit()
    
    flash('Xóa sáng kiến thành công!', 'success')
    return redirect(url_for('main.sangkien_list'))

# ==================== THỐNG KÊ ====================
@main_bp.route('/thongke')
@login_required
def thongke():
    """Trang thống kê tổng hợp với biểu đồ nâng cao"""
    # Lấy tham số lọc
    don_vi_filter = request.args.get('don_vi', '')
    nam_filter = request.args.get('nam', '')
    type_filter = request.args.get('type', 'all')
    
    # Năm hiện tại và năm trước
    current_year = datetime.now().year
    if nam_filter:
        current_year = int(nam_filter)
    previous_year = current_year - 1
    
    # ========== THỐNG KÊ ĐỀ TÀI THEO ĐƠN VỊ ==========
    query = db.session.query(
        DeTai.don_vi,
        func.count(DeTai.id).label('total')
    )
    if don_vi_filter:
        query = query.filter(DeTai.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', DeTai.ngay_bat_dau) == int(nam_filter))
    don_vi_stats = [{'don_vi': dv, 'total': total} for dv, total in query.group_by(DeTai.don_vi).all()]
    
    # ========== THỐNG KÊ SÁNG KIẾN THEO ĐƠN VỊ ==========
    query = db.session.query(
        SangKien.don_vi,
        func.count(SangKien.id).label('total')
    )
    if don_vi_filter:
        query = query.filter(SangKien.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', SangKien.ngay_bat_dau) == int(nam_filter))
    sangkien_stats = [{'don_vi': dv, 'total': total} for dv, total in query.group_by(SangKien.don_vi).all()]
    
    # ========== THỐNG KÊ TÀI LIỆU THEO LOẠI ==========
    query = db.session.query(
        TaiLieu.loai,
        func.count(TaiLieu.id).label('total')
    )
    if don_vi_filter:
        query = query.filter(TaiLieu.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', TaiLieu.ngay_bat_dau) == int(nam_filter))
    tailieu_stats = [{'loai': loai, 'total': total} for loai, total in query.group_by(TaiLieu.loai).all()]
    
    # ========== THỐNG KÊ TIỀM LỰC THEO ĐƠN VỊ ==========
    query = db.session.query(
        TiemLuc.don_vi,
        func.count(TiemLuc.id).label('total')
    )
    if don_vi_filter:
        query = query.filter(TiemLuc.don_vi == don_vi_filter)
    tiemluc_stats = [{'don_vi': dv, 'total': total} for dv, total in query.group_by(TiemLuc.don_vi).all()]
    
    # ========== THỐNG KÊ TRẠNG THÁI ĐỀ TÀI ==========
    query = db.session.query(
        DeTai.trang_thai,
        func.count(DeTai.id).label('total')
    )
    if don_vi_filter:
        query = query.filter(DeTai.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', DeTai.ngay_bat_dau) == int(nam_filter))
    detai_trangthai = [{'trang_thai': tt, 'total': total} for tt, total in query.group_by(DeTai.trang_thai).all()]
    
    # ========== THỐNG KÊ THEO THỜI GIAN (5 NĂM GẦN NHẤT) ==========
    detai_by_year = []
    sangkien_by_year = []
    tailieu_by_year = []
    
    for year in range(current_year - 4, current_year + 1):
        # Đề tài
        query = db.session.query(func.count(DeTai.id))
        if don_vi_filter:
            query = query.filter(DeTai.don_vi == don_vi_filter)
        query = query.filter(func.extract('year', DeTai.ngay_bat_dau) == year)
        detai_count = query.scalar() or 0
        detai_by_year.append({'year': year, 'count': detai_count})
        
        # Sáng kiến
        query = db.session.query(func.count(SangKien.id))
        if don_vi_filter:
            query = query.filter(SangKien.don_vi == don_vi_filter)
        query = query.filter(func.extract('year', SangKien.ngay_bat_dau) == year)
        sangkien_count = query.scalar() or 0
        sangkien_by_year.append({'year': year, 'count': sangkien_count})
        
        # Tài liệu
        query = db.session.query(func.count(TaiLieu.id))
        if don_vi_filter:
            query = query.filter(TaiLieu.don_vi == don_vi_filter)
        query = query.filter(func.extract('year', TaiLieu.ngay_bat_dau) == year)
        tailieu_count = query.scalar() or 0
        tailieu_by_year.append({'year': year, 'count': tailieu_count})
    
    # ========== THỐNG KÊ THEO CẤP ĐỘ ==========
    query = db.session.query(
        DeTai.cap_do,
        func.count(DeTai.id).label('total')
    )
    if don_vi_filter:
        query = query.filter(DeTai.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', DeTai.ngay_bat_dau) == int(nam_filter))
    detai_by_capdo = [{'cap_do': cd or 'Chưa xác định', 'total': total} 
                      for cd, total in query.group_by(DeTai.cap_do).all()]
    
    # ========== THỐNG KÊ HỌC VỊ/HỌC HÀM ==========
    hocvi_stats = db.session.query(
        TiemLuc.hoc_vi,
        func.count(TiemLuc.id).label('total')
    ).filter(TiemLuc.hoc_vi.isnot(None))
    if don_vi_filter:
        hocvi_stats = hocvi_stats.filter(TiemLuc.don_vi == don_vi_filter)
    hocvi_stats = [{'hoc_vi': hv or 'Chưa có', 'total': total} 
                   for hv, total in hocvi_stats.group_by(TiemLuc.hoc_vi).all()]
    
    hocham_stats = db.session.query(
        TiemLuc.hoc_ham,
        func.count(TiemLuc.id).label('total')
    ).filter(TiemLuc.hoc_ham.isnot(None))
    if don_vi_filter:
        hocham_stats = hocham_stats.filter(TiemLuc.don_vi == don_vi_filter)
    hocham_stats = [{'hoc_ham': hh or 'Chưa có', 'total': total} 
                    for hh, total in hocham_stats.group_by(TiemLuc.hoc_ham).all()]
    
    # ========== MỚI: SO SÁNH NĂM HIỆN TẠI VỚI NĂM TRƯỚC ==========
    # Đề tài năm hiện tại
    detai_current = db.session.query(func.count(DeTai.id)).filter(
        func.extract('year', DeTai.ngay_bat_dau) == current_year
    )
    if don_vi_filter:
        detai_current = detai_current.filter(DeTai.don_vi == don_vi_filter)
    detai_current_count = detai_current.scalar() or 0
    
    # Đề tài năm trước
    detai_previous = db.session.query(func.count(DeTai.id)).filter(
        func.extract('year', DeTai.ngay_bat_dau) == previous_year
    )
    if don_vi_filter:
        detai_previous = detai_previous.filter(DeTai.don_vi == don_vi_filter)
    detai_previous_count = detai_previous.scalar() or 0
    
    # Tính % tăng/giảm đề tài
    detai_change_percent = 0
    if detai_previous_count > 0:
        detai_change_percent = ((detai_current_count - detai_previous_count) / detai_previous_count) * 100
    
    # Sáng kiến năm hiện tại
    sangkien_current = db.session.query(func.count(SangKien.id)).filter(
        func.extract('year', SangKien.ngay_bat_dau) == current_year
    )
    if don_vi_filter:
        sangkien_current = sangkien_current.filter(SangKien.don_vi == don_vi_filter)
    sangkien_current_count = sangkien_current.scalar() or 0
    
    # Sáng kiến năm trước
    sangkien_previous = db.session.query(func.count(SangKien.id)).filter(
        func.extract('year', SangKien.ngay_bat_dau) == previous_year
    )
    if don_vi_filter:
        sangkien_previous = sangkien_previous.filter(SangKien.don_vi == don_vi_filter)
    sangkien_previous_count = sangkien_previous.scalar() or 0
    
    # Tính % tăng/giảm sáng kiến
    sangkien_change_percent = 0
    if sangkien_previous_count > 0:
        sangkien_change_percent = ((sangkien_current_count - sangkien_previous_count) / sangkien_previous_count) * 100
    
    # ========== MỚI: TOP 5 ĐƠN VỊ XUẤT SẮC ==========
    # Top đề tài
    top_detai = db.session.query(
        DeTai.don_vi,
        func.count(DeTai.id).label('total')
    )
    if nam_filter:
        top_detai = top_detai.filter(func.extract('year', DeTai.ngay_bat_dau) == int(nam_filter))
    top_detai = top_detai.group_by(DeTai.don_vi).order_by(func.count(DeTai.id).desc()).limit(5).all()
    top_detai_list = [{'don_vi': dv or 'Chưa xác định', 'total': total} for dv, total in top_detai]
    
    # Top sáng kiến
    top_sangkien = db.session.query(
        SangKien.don_vi,
        func.count(SangKien.id).label('total')
    )
    if nam_filter:
        top_sangkien = top_sangkien.filter(func.extract('year', SangKien.ngay_bat_dau) == int(nam_filter))
    top_sangkien = top_sangkien.group_by(SangKien.don_vi).order_by(func.count(SangKien.id).desc()).limit(5).all()
    top_sangkien_list = [{'don_vi': dv or 'Chưa xác định', 'total': total} for dv, total in top_sangkien]
    
    # ========== MỚI: THỐNG KÊ THEO QUÝ (NĂM HIỆN TẠI) ==========
    detai_by_quarter = []
    sangkien_by_quarter = []
    
    for quarter in range(1, 5):
        # Tháng bắt đầu và kết thúc của quý
        start_month = (quarter - 1) * 3 + 1
        end_month = quarter * 3
        
        # Đề tài theo quý
        query_detai = db.session.query(func.count(DeTai.id)).filter(
            func.extract('year', DeTai.ngay_bat_dau) == current_year,
            func.extract('month', DeTai.ngay_bat_dau) >= start_month,
            func.extract('month', DeTai.ngay_bat_dau) <= end_month
        )
        if don_vi_filter:
            query_detai = query_detai.filter(DeTai.don_vi == don_vi_filter)
        detai_count = query_detai.scalar() or 0
        detai_by_quarter.append({'quarter': f'Q{quarter}', 'count': detai_count})
        
        # Sáng kiến theo quý
        query_sangkien = db.session.query(func.count(SangKien.id)).filter(
            func.extract('year', SangKien.ngay_bat_dau) == current_year,
            func.extract('month', SangKien.ngay_bat_dau) >= start_month,
            func.extract('month', SangKien.ngay_bat_dau) <= end_month
        )
        if don_vi_filter:
            query_sangkien = query_sangkien.filter(SangKien.don_vi == don_vi_filter)
        sangkien_count = query_sangkien.scalar() or 0
        sangkien_by_quarter.append({'quarter': f'Q{quarter}', 'count': sangkien_count})
    
    # ========== MỚI: THỐNG KÊ TỶ LỆ HOÀN THÀNH ==========
    # Đề tài hoàn thành
    detai_completed = db.session.query(func.count(DeTai.id)).filter(
        DeTai.trang_thai == 'Hoàn thành'
    )
    if don_vi_filter:
        detai_completed = detai_completed.filter(DeTai.don_vi == don_vi_filter)
    if nam_filter:
        detai_completed = detai_completed.filter(func.extract('year', DeTai.ngay_bat_dau) == int(nam_filter))
    detai_completed_count = detai_completed.scalar() or 0
    
    # Tổng đề tài
    total_detai_query = db.session.query(func.count(DeTai.id))
    if don_vi_filter:
        total_detai_query = total_detai_query.filter(DeTai.don_vi == don_vi_filter)
    if nam_filter:
        total_detai_query = total_detai_query.filter(func.extract('year', DeTai.ngay_bat_dau) == int(nam_filter))
    total_detai_count = total_detai_query.scalar() or 1
    
    # Tỷ lệ hoàn thành
    completion_rate = (detai_completed_count / total_detai_count * 100) if total_detai_count > 0 else 0
    
    # ========== TỔNG HỢP ==========
    total_detai = sum(s['total'] for s in don_vi_stats)
    total_sangkien = sum(s['total'] for s in sangkien_stats)
    total_tailieu = sum(s['total'] for s in tailieu_stats)
    total_tiemluc = sum(s['total'] for s in tiemluc_stats)
    
    # ========== MỚI: TOP 10 NGƯỜI THAM GIA NHIỀU ĐỀ TÀI (Từ bảng ban_de_tai) ==========
    # Query từ bảng BanDeTai kết hợp với DeTai để lọc theo đơn vị và năm
    top_detai_participants = db.session.query(
        BanDeTai.ho_ten,
        func.count(BanDeTai.id).label('total')
    ).join(DeTai, BanDeTai.de_tai_id == DeTai.id)
    
    if don_vi_filter:
        top_detai_participants = top_detai_participants.filter(DeTai.don_vi == don_vi_filter)
    if nam_filter:
        top_detai_participants = top_detai_participants.filter(func.extract('year', DeTai.ngay_bat_dau) == int(nam_filter))
    
    top_detai_participants = top_detai_participants.group_by(BanDeTai.ho_ten)\
        .order_by(func.count(BanDeTai.id).desc())\
        .limit(10).all()
    top_detai_participants_list = [{'ho_ten': ht, 'total': total} for ht, total in top_detai_participants]
    
    # ========== MỚI: TOP 10 NGƯỜI THAM GIA NHIỀU SÁNG KIẾN (Từ bảng ban_sang_kien) ==========
    top_sangkien_participants = db.session.query(
        BanSangKien.ho_ten,
        func.count(BanSangKien.id).label('total')
    ).join(SangKien, BanSangKien.sang_kien_id == SangKien.id)
    
    if don_vi_filter:
        top_sangkien_participants = top_sangkien_participants.filter(SangKien.don_vi == don_vi_filter)
    if nam_filter:
        top_sangkien_participants = top_sangkien_participants.filter(func.extract('year', SangKien.ngay_bat_dau) == int(nam_filter))
    
    top_sangkien_participants = top_sangkien_participants.group_by(BanSangKien.ho_ten)\
        .order_by(func.count(BanSangKien.id).desc())\
        .limit(10).all()
    top_sangkien_participants_list = [{'ho_ten': ht, 'total': total} for ht, total in top_sangkien_participants]
    
    # ========== MỚI: TOP 10 NGƯỜI THAM GIA HỘI ĐỒNG NHIỀU ==========
    # Kết hợp từ tất cả các bảng hội đồng (đề tài, sáng kiến, tài liệu)
    from sqlalchemy import union_all
    
    # Query hội đồng đề cương đề tài
    hd_de_cuong_detai = db.session.query(
        HoiDongDeCuongDeTai.ho_ten.label('ho_ten'),
        HoiDongDeCuongDeTai.id.label('id'),
        DeTai.don_vi.label('don_vi'),
        DeTai.ngay_bat_dau.label('ngay_bat_dau')
    ).join(DeTai, HoiDongDeCuongDeTai.de_tai_id == DeTai.id)
    
    # Query hội đồng nghiệm thu đề tài
    hd_nghiem_thu_detai = db.session.query(
        HoiDongNghiemThuDeTai.ho_ten.label('ho_ten'),
        HoiDongNghiemThuDeTai.id.label('id'),
        DeTai.don_vi.label('don_vi'),
        DeTai.ngay_bat_dau.label('ngay_bat_dau')
    ).join(DeTai, HoiDongNghiemThuDeTai.de_tai_id == DeTai.id)
    
    # Query hội đồng đề cương sáng kiến
    hd_de_cuong_sangkien = db.session.query(
        HoiDongDeCuongSangKien.ho_ten.label('ho_ten'),
        HoiDongDeCuongSangKien.id.label('id'),
        SangKien.don_vi.label('don_vi'),
        SangKien.ngay_bat_dau.label('ngay_bat_dau')
    ).join(SangKien, HoiDongDeCuongSangKien.sang_kien_id == SangKien.id)
    
    # Query hội đồng nghiệm thu sáng kiến
    hd_nghiem_thu_sangkien = db.session.query(
        HoiDongNghiemThuSangKien.ho_ten.label('ho_ten'),
        HoiDongNghiemThuSangKien.id.label('id'),
        SangKien.don_vi.label('don_vi'),
        SangKien.ngay_bat_dau.label('ngay_bat_dau')
    ).join(SangKien, HoiDongNghiemThuSangKien.sang_kien_id == SangKien.id)
    
    # Query hội đồng đề cương tài liệu
    hd_de_cuong_tailieu = db.session.query(
        HoiDongDeCuongTaiLieu.ho_ten.label('ho_ten'),
        HoiDongDeCuongTaiLieu.id.label('id'),
        TaiLieu.don_vi.label('don_vi'),
        TaiLieu.ngay_bat_dau.label('ngay_bat_dau')
    ).join(TaiLieu, HoiDongDeCuongTaiLieu.tai_lieu_id == TaiLieu.id)
    
    # Query hội đồng nghiệm thu tài liệu
    hd_nghiem_thu_tailieu = db.session.query(
        HoiDongNghiemThuTaiLieu.ho_ten.label('ho_ten'),
        HoiDongNghiemThuTaiLieu.id.label('id'),
        TaiLieu.don_vi.label('don_vi'),
        TaiLieu.ngay_bat_dau.label('ngay_bat_dau')
    ).join(TaiLieu, HoiDongNghiemThuTaiLieu.tai_lieu_id == TaiLieu.id)
    
    # Kết hợp tất cả các query
    combined_hoidong = union_all(
        hd_de_cuong_detai,
        hd_nghiem_thu_detai,
        hd_de_cuong_sangkien,
        hd_nghiem_thu_sangkien,
        hd_de_cuong_tailieu,
        hd_nghiem_thu_tailieu
    ).alias('combined_hoidong')
    
    # Query top 10 từ kết quả kết hợp
    top_hoidong_query = db.session.query(
        combined_hoidong.c.ho_ten,
        func.count(combined_hoidong.c.id).label('total')
    )
    
    # Áp dụng filter
    if don_vi_filter:
        top_hoidong_query = top_hoidong_query.filter(combined_hoidong.c.don_vi == don_vi_filter)
    if nam_filter:
        top_hoidong_query = top_hoidong_query.filter(func.extract('year', combined_hoidong.c.ngay_bat_dau) == int(nam_filter))
    
    top_hoidong_participants = top_hoidong_query.group_by(combined_hoidong.c.ho_ten)\
        .order_by(func.count(combined_hoidong.c.id).desc())\
        .limit(10).all()
    
    top_hoidong_participants_list = [{'ho_ten': ht, 'total': total} for ht, total in top_hoidong_participants]
    
    # ========== MỚI: TOP 10 NGƯỜI THAM GIA NHIỀU TÀI LIỆU (Từ bảng ban_tai_lieu) ==========
    top_tailieu_participants = db.session.query(
        BanTaiLieu.ho_ten,
        func.count(BanTaiLieu.id).label('total')
    ).join(TaiLieu, BanTaiLieu.tai_lieu_id == TaiLieu.id)
    
    if don_vi_filter:
        top_tailieu_participants = top_tailieu_participants.filter(TaiLieu.don_vi == don_vi_filter)
    if nam_filter:
        top_tailieu_participants = top_tailieu_participants.filter(func.extract('year', TaiLieu.ngay_bat_dau) == int(nam_filter))
    
    top_tailieu_participants = top_tailieu_participants.group_by(BanTaiLieu.ho_ten)\
        .order_by(func.count(BanTaiLieu.id).desc())\
        .limit(10).all()
    top_tailieu_participants_list = [{'ho_ten': ht, 'total': total} for ht, total in top_tailieu_participants]
    
    # ========== MỚI: THỐNG KÊ TIỀM LỰC THEO CHỨC DANH VÀ ĐƠN VỊ ==========
    tiemluc_chuc_danh_stats = db.session.query(
        TiemLuc.don_vi,
        TiemLuc.chuc_danh,
        func.count(TiemLuc.id).label('total')
    )
    if don_vi_filter:
        tiemluc_chuc_danh_stats = tiemluc_chuc_danh_stats.filter(TiemLuc.don_vi == don_vi_filter)
    tiemluc_chuc_danh_stats = tiemluc_chuc_danh_stats.filter(
        TiemLuc.chuc_danh.isnot(None),
        TiemLuc.don_vi.isnot(None)
    ).group_by(TiemLuc.don_vi, TiemLuc.chuc_danh)\
     .order_by(TiemLuc.don_vi, func.count(TiemLuc.id).desc()).all()
    
    # Chuyển đổi thành cấu trúc dễ sử dụng trong template
    tiemluc_by_donvi_chucdanh = {}
    for don_vi, chuc_danh, total in tiemluc_chuc_danh_stats:
        if don_vi not in tiemluc_by_donvi_chucdanh:
            tiemluc_by_donvi_chucdanh[don_vi] = []
        tiemluc_by_donvi_chucdanh[don_vi].append({
            'chuc_danh': chuc_danh,
            'total': total
        })
    
    return render_template('thongke.html',
                         don_vi_stats=don_vi_stats,
                         sangkien_stats=sangkien_stats,
                         tailieu_stats=tailieu_stats,
                         tiemluc_stats=tiemluc_stats,
                         detai_trangthai=detai_trangthai,
                         detai_by_year=detai_by_year,
                         sangkien_by_year=sangkien_by_year,
                         tailieu_by_year=tailieu_by_year,
                         detai_by_capdo=detai_by_capdo,
                         hocvi_stats=hocvi_stats,
                         hocham_stats=hocham_stats,
                         total_detai=total_detai,
                         total_sangkien=total_sangkien,
                         total_tailieu=total_tailieu,
                         total_tiemluc=total_tiemluc,
                         don_vi_list=DON_VI_LIST,
                         # Thống kê mới
                         detai_current_count=detai_current_count,
                         detai_previous_count=detai_previous_count,
                         detai_change_percent=detai_change_percent,
                         sangkien_current_count=sangkien_current_count,
                         sangkien_previous_count=sangkien_previous_count,
                         sangkien_change_percent=sangkien_change_percent,
                         top_detai_list=top_detai_list,
                         top_sangkien_list=top_sangkien_list,
                         detai_by_quarter=detai_by_quarter,
                         sangkien_by_quarter=sangkien_by_quarter,
                         completion_rate=completion_rate,
                         current_year=current_year,
                         previous_year=previous_year,
                         # Thống kê mới - Top participants
                         top_detai_participants=top_detai_participants_list,
                         top_sangkien_participants=top_sangkien_participants_list,
                         top_hoidong_participants=top_hoidong_participants_list,
                         top_tailieu_participants=top_tailieu_participants_list,
                         # Thống kê tiềm lực theo chức danh và đơn vị
                         tiemluc_by_donvi_chucdanh=tiemluc_by_donvi_chucdanh)

@main_bp.route('/thongke/export-excel')
@login_required
def thongke_export_excel():
    """Xuất thống kê ra Excel"""
    # Lấy tham số lọc
    don_vi_filter = request.args.get('don_vi', '')
    nam_filter = request.args.get('nam', '')
    
    # Tạo workbook
    wb = Workbook()
    wb.remove(wb.active)  # Xóa sheet mặc định
    
    # Định dạng tiêu đề
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Sheet 1: Thống kê đề tài theo đơn vị
    ws1 = wb.create_sheet("Đề tài theo đơn vị")
    ws1.append(["STT", "Đơn vị", "Số lượng"])
    
    # Định dạng header
    for cell in ws1[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    query = db.session.query(DeTai.don_vi, func.count(DeTai.id).label('total'))
    if don_vi_filter:
        query = query.filter(DeTai.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', DeTai.ngay_bat_dau) == int(nam_filter))
    detai_stats = query.group_by(DeTai.don_vi).all()
    
    for idx, (don_vi, total) in enumerate(detai_stats, 1):
        ws1.append([idx, don_vi or "Chưa xác định", total])
    
    ws1.column_dimensions['A'].width = 10
    ws1.column_dimensions['B'].width = 40
    ws1.column_dimensions['C'].width = 15
    
    # Sheet 2: Thống kê sáng kiến theo đơn vị
    ws2 = wb.create_sheet("Sáng kiến theo đơn vị")
    ws2.append(["STT", "Đơn vị", "Số lượng"])
    
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    query = db.session.query(SangKien.don_vi, func.count(SangKien.id).label('total'))
    if don_vi_filter:
        query = query.filter(SangKien.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', SangKien.ngay_bat_dau) == int(nam_filter))
    sangkien_stats = query.group_by(SangKien.don_vi).all()
    
    for idx, (don_vi, total) in enumerate(sangkien_stats, 1):
        ws2.append([idx, don_vi or "Chưa xác định", total])
    
    ws2.column_dimensions['A'].width = 10
    ws2.column_dimensions['B'].width = 40
    ws2.column_dimensions['C'].width = 15
    
    # Sheet 3: Thống kê tài liệu theo đơn vị
    ws3 = wb.create_sheet("Tài liệu theo đơn vị")
    ws3.append(["STT", "Đơn vị", "Số lượng"])
    
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    query = db.session.query(TaiLieu.don_vi, func.count(TaiLieu.id).label('total'))
    if don_vi_filter:
        query = query.filter(TaiLieu.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', TaiLieu.ngay_bat_dau) == int(nam_filter))
    tailieu_stats = query.group_by(TaiLieu.don_vi).all()
    
    for idx, (don_vi, total) in enumerate(tailieu_stats, 1):
        ws3.append([idx, don_vi or "Chưa xác định", total])
    
    ws3.column_dimensions['A'].width = 10
    ws3.column_dimensions['B'].width = 40
    ws3.column_dimensions['C'].width = 15
    
    # Sheet 4: Thống kê tiềm lực theo đơn vị
    ws4 = wb.create_sheet("Tiềm lực theo đơn vị")
    ws4.append(["STT", "Đơn vị", "Số lượng"])
    
    for cell in ws4[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    query = db.session.query(TiemLuc.don_vi, func.count(TiemLuc.id).label('total'))
    if don_vi_filter:
        query = query.filter(TiemLuc.don_vi == don_vi_filter)
    tiemluc_stats = query.group_by(TiemLuc.don_vi).all()
    
    for idx, (don_vi, total) in enumerate(tiemluc_stats, 1):
        ws4.append([idx, don_vi or "Chưa xác định", total])
    
    ws4.column_dimensions['A'].width = 10
    ws4.column_dimensions['B'].width = 40
    ws4.column_dimensions['C'].width = 15
    
    # Sheet 5: Tổng hợp
    ws5 = wb.create_sheet("Tổng hợp", 0)
    ws5.append(["THỐNG KÊ KHOA HỌC"])
    ws5.merge_cells('A1:C1')
    ws5['A1'].font = Font(bold=True, size=16)
    ws5['A1'].alignment = Alignment(horizontal="center")
    
    ws5.append([])
    ws5.append(["Loại", "Số lượng"])
    
    for cell in ws5[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    total_detai = DeTai.query.count() if not don_vi_filter and not nam_filter else len(detai_stats)
    total_sangkien = SangKien.query.count() if not don_vi_filter and not nam_filter else len(sangkien_stats)
    total_tailieu = TaiLieu.query.count() if not don_vi_filter and not nam_filter else len(tailieu_stats)
    total_tiemluc = TiemLuc.query.count() if not don_vi_filter else len(tiemluc_stats)
    
    ws5.append(["Đề tài", sum(total for _, total in detai_stats)])
    ws5.append(["Sáng kiến", sum(total for _, total in sangkien_stats)])
    ws5.append(["Tài liệu", sum(total for _, total in tailieu_stats)])
    ws5.append(["Tiềm lực", sum(total for _, total in tiemluc_stats)])
    
    ws5.column_dimensions['A'].width = 30
    ws5.column_dimensions['B'].width = 20
    
    # Lưu file vào memory
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Tạo tên file
    filename = f"thongke_khoa_hoc"
    if don_vi_filter:
        filename += f"_{don_vi_filter}"
    if nam_filter:
        filename += f"_{nam_filter}"
    filename += ".xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

@main_bp.route('/thongke/export-word')
@login_required  
def thongke_export_word():
    """Xuất thống kê ra Word"""
    from flask import send_file
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import io
    
    # Lấy tham số lọc
    don_vi_filter = request.args.get('don_vi', '')
    nam_filter = request.args.get('nam', '')
    
    # Tạo document
    doc = Document()
    
    # Tiêu đề
    heading = doc.add_heading('BÁO CÁO THỐNG KÊ HOẠT ĐỘNG KHOA HỌC', 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Thông tin bộ lọc
    if don_vi_filter or nam_filter:
        filter_info = doc.add_paragraph()
        filter_info.add_run('Bộ lọc: ').bold = True
        if don_vi_filter:
            filter_info.add_run(f'Đơn vị: {don_vi_filter}  ')
        if nam_filter:
            filter_info.add_run(f'Năm: {nam_filter}')
        filter_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # 1. Thống kê đề tài
    doc.add_heading('I. THỐNG KÊ ĐỀ TÀI THEO ĐƠN VỊ', 1)
    query = db.session.query(DeTai.don_vi, func.count(DeTai.id).label('total'))
    if don_vi_filter:
        query = query.filter(DeTai.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', DeTai.ngay_bat_dau) == int(nam_filter))
    detai_stats = query.group_by(DeTai.don_vi).all()
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'STT'
    hdr_cells[1].text = 'Đơn vị'
    hdr_cells[2].text = 'Số lượng'
    
    for idx, (don_vi, total) in enumerate(detai_stats, 1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = don_vi or "Chưa xác định"
        row_cells[2].text = str(total)
    
    doc.add_paragraph()
    
    # 2. Thống kê sáng kiến
    doc.add_heading('II. THỐNG KÊ SÁNG KIẾN THEO ĐƠN VỊ', 1)
    query = db.session.query(SangKien.don_vi, func.count(SangKien.id).label('total'))
    if don_vi_filter:
        query = query.filter(SangKien.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', SangKien.ngay_bat_dau) == int(nam_filter))
    sangkien_stats = query.group_by(SangKien.don_vi).all()
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'STT'
    hdr_cells[1].text = 'Đơn vị'
    hdr_cells[2].text = 'Số lượng'
    
    for idx, (don_vi, total) in enumerate(sangkien_stats, 1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = don_vi or "Chưa xác định"
        row_cells[2].text = str(total)
    
    doc.add_paragraph()
    
    # 3. Thống kê tài liệu
    doc.add_heading('III. THỐNG KÊ TÀI LIỆU THEO LOẠI', 1)
    query = db.session.query(TaiLieu.loai, func.count(TaiLieu.id).label('total'))
    if don_vi_filter:
        query = query.filter(TaiLieu.don_vi == don_vi_filter)
    if nam_filter:
        query = query.filter(func.extract('year', TaiLieu.ngay_bat_dau) == int(nam_filter))
    tailieu_stats = query.group_by(TaiLieu.loai).all()
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'STT'
    hdr_cells[1].text = 'Loại tài liệu'
    hdr_cells[2].text = 'Số lượng'
    
    for idx, (loai, total) in enumerate(tailieu_stats, 1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = loai or "Chưa phân loại"
        row_cells[2].text = str(total)
    
    doc.add_paragraph()
    
    # 4. Thống kê tiềm lực
    doc.add_heading('IV. THỐNG KÊ TIỀM LỰC KHOA HỌC THEO ĐƠN VỊ', 1)
    query = db.session.query(TiemLuc.don_vi, func.count(TiemLuc.id).label('total'))
    if don_vi_filter:
        query = query.filter(TiemLuc.don_vi == don_vi_filter)
    tiemluc_stats = query.group_by(TiemLuc.don_vi).all()
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'STT'
    hdr_cells[1].text = 'Đơn vị'
    hdr_cells[2].text = 'Số lượng'
    
    for idx, (don_vi, total) in enumerate(tiemluc_stats, 1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = don_vi or "Chưa xác định"
        row_cells[2].text = str(total)
    
    # Lưu vào memory
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    # Tạo tên file
    filename = f"baocao_thongke_khoa_hoc"
    if don_vi_filter:
        filename += f"_{don_vi_filter.replace(' ', '_')}"
    if nam_filter:
        filename += f"_{nam_filter}"
    filename += ".docx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=filename
    )



# ==================== DOWNLOAD FILE ====================
@main_bp.route('/download/<path:filepath>')
@login_required
def download_file(filepath):
    """Download file đã upload"""
    try:
        # Đường dẫn đầy đủ đến file
        full_path = os.path.join(filepath)
        
        # Kiểm tra file có tồn tại không
        if not os.path.exists(full_path):
            flash('File không tồn tại!', 'danger')
            return redirect(request.referrer or url_for('main.dashboard'))
        
        # Lấy tên file gốc
        filename = os.path.basename(full_path)
        
        # Gửi file
        return send_file(full_path, as_attachment=True, download_name=filename)
    except Exception as e:
        flash(f'Lỗi khi tải file: {str(e)}', 'danger')
        return redirect(request.referrer or url_for('main.dashboard'))
# === DETAIL ROUTES ===

@main_bp.route('/detai/detail/<int:id>')
@login_required
def detai_detail(id):
    detai = DeTai.query.get_or_404(id)
    hoi_dong_de_cuong_members = HoiDongDeCuongDeTai.query.filter_by(de_tai_id=id).all()
    ban_members = BanDeTai.query.filter_by(de_tai_id=id).all()

    hoi_dong_members = HoiDongNghiemThuDeTai.query.filter_by(de_tai_id=id).all()
    return render_template('detai/detail.html', 
                         detai=detai,
                         ban_members=ban_members,
                         hoi_dong_members=hoi_dong_members,
                         hoi_dong_de_cuong_members=hoi_dong_de_cuong_members)
                         


@main_bp.route('/sangkien/detail/<int:id>')
@login_required
def sangkien_detail(id):
    sangkien = SangKien.query.get_or_404(id)
    hoi_dong_de_cuong_members = HoiDongDeCuongSangKien.query.filter_by(sang_kien_id=id).all()
    ban_members = BanSangKien.query.filter_by(sang_kien_id=id).all()
    hoi_dong_members = HoiDongNghiemThuSangKien.query.filter_by(sang_kien_id=id).all()

    return render_template('sangkien/detail.html', sangkien=sangkien,
                         ban_members=ban_members,
                         hoi_dong_members=hoi_dong_members,
                         hoi_dong_de_cuong_members=hoi_dong_de_cuong_members)


@main_bp.route('/tiemluc/detail/<int:id>')
@login_required
def tiemluc_detail(id):
    tiemluc = TiemLuc.query.get_or_404(id)
    
    # Lấy danh sách công trình nghiên cứu của cán bộ này
    # detai_list = DeTai.query.filter(
    #     (DeTai.chu_nhiem.like(f'%{tiemluc.ho_ten}%')) | 
    #     (DeTai.ban_thanh_vien.like(f'%{tiemluc.ho_ten}%'))
    # ).all()
    
    # sangkien_list = SangKien.query.filter(
    #     (SangKien.chu_nhiem.like(f'%{tiemluc.ho_ten}%')) | 
    #     (SangKien.ban_thanh_vien.like(f'%{tiemluc.ho_ten}%'))
    # ).all()
    
    # tailieu_list = TaiLieu.query.filter(
    #     TaiLieu.tac_gia.like(f'%{tiemluc.ho_ten}%')
    # ).all()
    
    return render_template('tiemluc/detail.html', 
                         tiemluc=tiemluc,
                         detai_list=detai_list,
                         sangkien_list=sangkien_list,
                         tailieu_list=tailieu_list)

@main_bp.route('/timkiem', methods=['GET'])
@login_required
def timkiem():
    search_query = request.args.get('q', '').strip()
    search_type = request.args.get('type', 'all')
    
    results = {
        'detai': [],
        'sangkien': [],
        'tailieu': [],
        'tiemluc': [],
        'file_matches': []  # Kết quả tìm trong file
    }
    
    person_stats = []
    file_search_query = ''
    
    if search_query:
        # ============= TÌM KIẾM THEO TÊN NGƯỜI =============
        if search_type == 'person':
            # Tìm trong các bảng liên quan
            person_data = {}
            
            # 1. Tìm trong Ban đề tài
            ban_detai = BanDeTai.query.filter(
                BanDeTai.ho_ten.ilike(f'%{search_query}%')
            ).all()
            for item in ban_detai:
                name = item.ho_ten
                if name not in person_data:
                    person_data[name] = {
                        'name': name,
                        'detai_list': [],
                        'sangkien_list': [],
                        'tailieu_list': [],
                        'hoi_dong_list': []
                    }
                detai = DeTai.query.get(item.de_tai_id)
                if detai:
                    person_data[name]['detai_list'].append({
                        'id': detai.id,
                        'ten_de_tai': detai.ten_de_tai,
                        'vai_tro': item.vai_tro
                    })
            
            # 2. Tìm trong Ban sáng kiến
            ban_sangkien = BanSangKien.query.filter(
                BanSangKien.ho_ten.ilike(f'%{search_query}%')
            ).all()
            for item in ban_sangkien:
                name = item.ho_ten
                if name not in person_data:
                    person_data[name] = {
                        'name': name,
                        'detai_list': [],
                        'sangkien_list': [],
                        'tailieu_list': [],
                        'hoi_dong_list': []
                    }
                sangkien = SangKien.query.get(item.sang_kien_id)
                if sangkien:
                    person_data[name]['sangkien_list'].append({
                        'id': sangkien.id,
                        'ten_sang_kien': sangkien.ten_sang_kien,
                        'vai_tro': item.vai_tro
                    })
            
            # 3. Tìm trong Ban tài liệu
            ban_tailieu = BanTaiLieu.query.filter(
                BanTaiLieu.ho_ten.ilike(f'%{search_query}%')
            ).all()
            for item in ban_tailieu:
                name = item.ho_ten
                if name not in person_data:
                    person_data[name] = {
                        'name': name,
                        'detai_list': [],
                        'sangkien_list': [],
                        'tailieu_list': [],
                        'hoi_dong_list': []
                    }
                tailieu = TaiLieu.query.get(item.tai_lieu_id)
                if tailieu:
                    person_data[name]['tailieu_list'].append({
                        'id': tailieu.id,
                        'ten_tai_lieu': tailieu.ten_tai_lieu,
                        'vai_tro': item.vai_tro
                    })
            
            # 4. Tìm trong Hội đồng đề cương đề tài
            hd_dc_detai = HoiDongDeCuongDeTai.query.filter(
                HoiDongDeCuongDeTai.ho_ten.ilike(f'%{search_query}%')
            ).all()
            for item in hd_dc_detai:
                name = item.ho_ten
                if name not in person_data:
                    person_data[name] = {
                        'name': name,
                        'detai_list': [],
                        'sangkien_list': [],
                        'tailieu_list': [],
                        'hoi_dong_list': []
                    }
                detai = DeTai.query.get(item.de_tai_id)
                if detai:
                    person_data[name]['hoi_dong_list'].append({
                        'loai_hoi_dong': 'HĐ Đề cương Đề tài',
                        'ten': detai.ten_de_tai,
                        'vai_tro': item.vai_tro,
                        'url': url_for('main.detai_detail', id=detai.id)
                    })
            
            # 5. Tìm trong Hội đồng nghiệm thu đề tài
            hd_nt_detai = HoiDongNghiemThuDeTai.query.filter(
                HoiDongNghiemThuDeTai.ho_ten.ilike(f'%{search_query}%')
            ).all()
            for item in hd_nt_detai:
                name = item.ho_ten
                if name not in person_data:
                    person_data[name] = {
                        'name': name,
                        'detai_list': [],
                        'sangkien_list': [],
                        'tailieu_list': [],
                        'hoi_dong_list': []
                    }
                detai = DeTai.query.get(item.de_tai_id)
                if detai:
                    person_data[name]['hoi_dong_list'].append({
                        'loai_hoi_dong': 'HĐ Nghiệm thu Đề tài',
                        'ten': detai.ten_de_tai,
                        'vai_tro': item.vai_tro,
                        'url': url_for('main.detai_detail', id=detai.id)
                    })
            
            # 6. Tìm trong Hội đồng đề cương sáng kiến
            hd_dc_sangkien = HoiDongDeCuongSangKien.query.filter(
                HoiDongDeCuongSangKien.ho_ten.ilike(f'%{search_query}%')
            ).all()
            for item in hd_dc_sangkien:
                name = item.ho_ten
                if name not in person_data:
                    person_data[name] = {
                        'name': name,
                        'detai_list': [],
                        'sangkien_list': [],
                        'tailieu_list': [],
                        'hoi_dong_list': []
                    }
                sangkien = SangKien.query.get(item.sang_kien_id)
                if sangkien:
                    person_data[name]['hoi_dong_list'].append({
                        'loai_hoi_dong': 'HĐ Đề cương Sáng kiến',
                        'ten': sangkien.ten_sang_kien,
                        'vai_tro': item.vai_tro,
                        'url': url_for('main.sangkien_detail', id=sangkien.id)
                    })
            
            # 7. Tìm trong Hội đồng nghiệm thu sáng kiến
            hd_nt_sangkien = HoiDongNghiemThuSangKien.query.filter(
                HoiDongNghiemThuSangKien.ho_ten.ilike(f'%{search_query}%')
            ).all()
            for item in hd_nt_sangkien:
                name = item.ho_ten
                if name not in person_data:
                    person_data[name] = {
                        'name': name,
                        'detai_list': [],
                        'sangkien_list': [],
                        'tailieu_list': [],
                        'hoi_dong_list': []
                    }
                sangkien = SangKien.query.get(item.sang_kien_id)
                if sangkien:
                    person_data[name]['hoi_dong_list'].append({
                        'loai_hoi_dong': 'HĐ Nghiệm thu Sáng kiến',
                        'ten': sangkien.ten_sang_kien,
                        'vai_tro': item.vai_tro,
                        'url': url_for('main.sangkien_detail', id=sangkien.id)
                    })
            
            # 8. Tìm trong Hội đồng đề cương tài liệu
            hd_dc_tailieu = HoiDongDeCuongTaiLieu.query.filter(
                HoiDongDeCuongTaiLieu.ho_ten.ilike(f'%{search_query}%')
            ).all()
            for item in hd_dc_tailieu:
                name = item.ho_ten
                if name not in person_data:
                    person_data[name] = {
                        'name': name,
                        'detai_list': [],
                        'sangkien_list': [],
                        'tailieu_list': [],
                        'hoi_dong_list': []
                    }
                tailieu = TaiLieu.query.get(item.tai_lieu_id)
                if tailieu:
                    person_data[name]['hoi_dong_list'].append({
                        'loai_hoi_dong': 'HĐ Đề cương Tài liệu',
                        'ten': tailieu.ten_tai_lieu,
                        'vai_tro': item.vai_tro,
                        'url': url_for('main.tailieu_detail', id=tailieu.id)
                    })
            
            # 9. Tìm trong Hội đồng nghiệm thu tài liệu
            hd_nt_tailieu = HoiDongNghiemThuTaiLieu.query.filter(
                HoiDongNghiemThuTaiLieu.ho_ten.ilike(f'%{search_query}%')
            ).all()
            for item in hd_nt_tailieu:
                name = item.ho_ten
                if name not in person_data:
                    person_data[name] = {
                        'name': name,
                        'detai_list': [],
                        'sangkien_list': [],
                        'tailieu_list': [],
                        'hoi_dong_list': []
                    }
                tailieu = TaiLieu.query.get(item.tai_lieu_id)
                if tailieu:
                    person_data[name]['hoi_dong_list'].append({
                        'loai_hoi_dong': 'HĐ Nghiệm thu Tài liệu',
                        'ten': tailieu.ten_tai_lieu,
                        'vai_tro': item.vai_tro,
                        'url': url_for('main.tailieu_detail', id=tailieu.id)
                    })
            
            # Tính toán thống kê
            for name, data in person_data.items():
                person_stats.append({
                    'name': name,
                    'detai_count': len(data['detai_list']),
                    'sangkien_count': len(data['sangkien_list']),
                    'tailieu_count': len(data['tailieu_list']),
                    'hoi_dong_count': len(data['hoi_dong_list']),
                    'total_participations': len(data['detai_list']) + len(data['sangkien_list']) + len(data['tailieu_list']) + len(data['hoi_dong_list']),
                    'detai_list': data['detai_list'],
                    'sangkien_list': data['sangkien_list'],
                    'tailieu_list': data['tailieu_list'],
                    'hoi_dong_list': data['hoi_dong_list']
                })
            
            # Sắp xếp theo tổng số lần tham gia
            person_stats.sort(key=lambda x: x['total_participations'], reverse=True)
        
        # ============= TÌM KIẾM TRONG FILE =============
        elif search_type == 'file_content':
            file_search_query = search_query
            upload_folder = os.path.join('static', 'uploads')
            
            # Tìm trong file đề tài
            for detai in DeTai.query.all():
                files_to_check = []
                if detai.file_sau_sua:
                    files_to_check.append(('Nội dung', detai.file_sau_sua))
                # if detai.bien_ban_de_cuong:
                #     files_to_check.append(('Biên bản họp', detai.bien_ban_de_cuong))
                # if detai.bien_ban_nghiem_thu:
                #     files_to_check.append(('Nghiệm thu', detai.bien_ban_nghiem_thu))
                
                for file_type, file_path in files_to_check:
                    full_path = os.path.join(upload_folder, file_path)
                    found, snippet = search_in_file(full_path, search_query)
                    if found:
                        results['file_matches'].append({
                            'type': 'detai',
                            'item': detai,
                            'file_type': file_type,
                            'file_path': file_path,
                            'snippet': snippet,
                            'title': detai.ten_de_tai,
                            'code': detai.ma_de_tai,
                            'url': url_for('main.detai_detail', id=detai.id)
                        })
            
            # Tìm trong file sáng kiến
            for sangkien in SangKien.query.all():
                files_to_check = []
                if sangkien.file_sau_sua:
                    files_to_check.append(('Hồ sơ sáng kiến', sangkien.file_sau_sua))
                if sangkien.bien_ban_de_cuong:
                    files_to_check.append(('Giấy chứng nhận', sangkien.bien_ban_de_cuong))
                if sangkien.bien_ban_nghiem_thu:
                    files_to_check.append(('BC Ứng dụng', sangkien.bien_ban_nghiem_thu))
                
                for file_type, file_path in files_to_check:
                    full_path = os.path.join(upload_folder, file_path)
                    found, snippet = search_in_file(full_path, search_query)
                    if found:
                        results['file_matches'].append({
                            'type': 'sangkien',
                            'item': sangkien,
                            'file_type': file_type,
                            'file_path': file_path,
                            'snippet': snippet,
                            'title': sangkien.ten_sang_kien,
                            'code': sangkien.ma_sang_kien,
                            'url': url_for('main.sangkien_detail', id=sangkien.id)
                        })
            
            # Tìm trong file tài liệu
            for tailieu in TaiLieu.query.all():
                files_to_check = []
                if tailieu.file_sau_sua:
                    files_to_check.append(('Tài liệu', tailieu.file_sau_sua))
                if tailieu.bien_ban_ban:
                    files_to_check.append(('BB Ban', tailieu.bien_ban_ban))
                if tailieu.bien_ban_de_cuong:
                    files_to_check.append(('BB Đề cương', tailieu.bien_ban_de_cuong))
                if tailieu.bien_ban_nghiem_thu:
                    files_to_check.append(('BB Nghiệm thu', tailieu.bien_ban_nghiem_thu))
                
                for file_type, file_path in files_to_check:
                    full_path = os.path.join(upload_folder, file_path)
                    found, snippet = search_in_file(full_path, search_query)
                    if found:
                        results['file_matches'].append({
                            'type': 'tailieu',
                            'item': tailieu,
                            'file_type': file_type,
                            'file_path': file_path,
                            'snippet': snippet,
                            'title': tailieu.ten_tai_lieu,
                            'code': tailieu.ma_tai_lieu,
                            'url': url_for('main.tailieu_detail', id=tailieu.id)
                        })
        
        # ============= TÌM KIẾM THÔNG THƯỜNG =============
        else:
            # Tìm kiếm thông thường trong database
            if search_type in ['all', 'detai']:
                results['detai'] = DeTai.query.filter(
                    db.or_(
                        DeTai.ma_de_tai.ilike(f'%{search_query}%'),
                        DeTai.ten_de_tai.ilike(f'%{search_query}%'),
                        DeTai.chu_nhiem.ilike(f'%{search_query}%'),
                    #     DeTai.ban_thanh_vien.ilike(f'%{search_query}%'),
                        DeTai.mo_ta.ilike(f'%{search_query}%'),
                        DeTai.don_vi.ilike(f'%{search_query}%')
                    )
                ).all()
            
            if search_type in ['all', 'sangkien']:
                results['sangkien'] = SangKien.query.filter(
                    db.or_(
                        SangKien.ma_sang_kien.ilike(f'%{search_query}%'),
                        SangKien.ten_sang_kien.ilike(f'%{search_query}%'),
                        SangKien.chu_nhiem.ilike(f'%{search_query}%'),
                    #   SangKien.ban_thanh_vien.ilike(f'%{search_query}%'),
                        SangKien.mo_ta.ilike(f'%{search_query}%'),
                        SangKien.don_vi.ilike(f'%{search_query}%')
                    )
                ).all()
            
            if search_type in ['all', 'tailieu']:
                results['tailieu'] = TaiLieu.query.filter(
                    db.or_(
                        TaiLieu.ma_tai_lieu.ilike(f'%{search_query}%'),
                        TaiLieu.ten_tai_lieu.ilike(f'%{search_query}%'),
                        TaiLieu.tac_gia.ilike(f'%{search_query}%'),
                        TaiLieu.mo_ta.ilike(f'%{search_query}%'),
                        TaiLieu.don_vi.ilike(f'%{search_query}%')
                    )
                ).all()
            
            if search_type in ['all', 'tiemluc']:
                results['tiemluc'] = TiemLuc.query.filter(
                    db.or_(
                        TiemLuc.ma_can_bo.ilike(f'%{search_query}%'),
                        TiemLuc.ho_ten.ilike(f'%{search_query}%'),
                        TiemLuc.don_vi.ilike(f'%{search_query}%'),
                        TiemLuc.chuyen_mon.ilike(f'%{search_query}%'),
                        TiemLuc.kinh_nghiem.ilike(f'%{search_query}%')
                    )
                ).all()
    
    return render_template('timkiem.html', 
                         search_query=search_query,
                         search_type=search_type,
                         results=results,
                         person_stats=person_stats,
                         file_search_query=file_search_query)